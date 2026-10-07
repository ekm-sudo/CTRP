"""What: checks the officers' workbook has its fixed sheets.  Why: officers rely on the layout."""
from openpyxl import load_workbook

from ctrp import core as M
from ctrp import report as R


def test_workbook_has_the_agreed_sheets(tmp_path, extract):
    out = R.write_workbook(M.score_snapshot(extract), tmp_path / 'scores.xlsx', 'test.csv')
    sheets = load_workbook(out, read_only=True).sheetnames
    assert sheets == ['Read me', 'Scores', 'Reasons', 'Parameter values', 'Band summary', 'Reason codes', 'Run info']
